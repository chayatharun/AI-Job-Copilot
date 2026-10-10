import os
import re
from datetime import date, datetime
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, jsonify, request
from flask_cors import CORS
from sqlalchemy import Date, DateTime, Integer, String, Text, create_engine, select, desc
from sqlalchemy.orm import DeclarativeBase, Mapped, mapped_column, sessionmaker
from werkzeug.utils import secure_filename

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document
except ImportError:
    Document = None

load_dotenv()
app = Flask(__name__)
CORS(app)
app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024
ALLOWED_EXTENSIONS = {"pdf", "docx"}

DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
if not DATABASE_URL:
    # Works immediately without MySQL; set DATABASE_URL in .env to use MySQL.
    DATABASE_URL = "sqlite:///jobcopilot_local.db"

engine_options = {"pool_pre_ping": True}
if DATABASE_URL.startswith("sqlite"):
    engine_options["connect_args"] = {"check_same_thread": False}
engine = create_engine(DATABASE_URL, **engine_options)
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)


class Base(DeclarativeBase):
    pass


class SavedAnalysis(Base):
    __tablename__ = "saved_job_analyses"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow, nullable=False)
    file_name: Mapped[str] = mapped_column(String(255), default="Resume", nullable=False)
    job_description: Mapped[str] = mapped_column(Text, nullable=False)
    match_score: Mapped[int] = mapped_column(Integer, nullable=False)
    resume_skills_json: Mapped[str] = mapped_column(Text, default="[]", nullable=False)
    job_skills_json: Mapped[str] = mapped_column(Text, default="[]", nullable=False)
    matching_skills_json: Mapped[str] = mapped_column(Text, default="[]", nullable=False)
    missing_skills_json: Mapped[str] = mapped_column(Text, default="[]", nullable=False)
    recommendations_json: Mapped[str] = mapped_column(Text, default="[]", nullable=False)
    extracted_character_count: Mapped[int] = mapped_column(Integer, default=0, nullable=False)
    detected_job_skill_count: Mapped[int] = mapped_column(Integer, default=0, nullable=False)
    analysis_method: Mapped[str] = mapped_column(String(100), default="Rule-based skill comparison", nullable=False)


class JobApplication(Base):
    __tablename__ = "job_applications"
    id: Mapped[int] = mapped_column(Integer, primary_key=True, autoincrement=True)
    job_title: Mapped[str] = mapped_column(String(255), nullable=False)
    company: Mapped[str] = mapped_column(String(255), nullable=False)
    job_url: Mapped[str | None] = mapped_column(Text, nullable=True)
    status: Mapped[str] = mapped_column(String(40), default="Saved", nullable=False)
    applied_date: Mapped[date | None] = mapped_column(Date, nullable=True)
    notes: Mapped[str | None] = mapped_column(Text, nullable=True)
    created_at: Mapped[datetime] = mapped_column(DateTime, default=datetime.utcnow, nullable=False)


Base.metadata.create_all(engine)


SKILL_LIBRARY = [
    "Python", "Java", "JavaScript", "TypeScript", "C", "C++", "SQL", "MySQL",
    "PostgreSQL", "SQLite", "HTML", "CSS", "React", "Node.js", "Express",
    "Flask", "FastAPI", "Spring Boot", "REST API", "Git", "GitHub", "Docker",
    "AWS", "Azure", "Machine Learning", "Deep Learning", "NLP", "Pandas",
    "NumPy", "scikit-learn", "TensorFlow", "PyTorch", "XGBoost", "Power BI",
    "Tableau", "Excel", "Selenium", "JUnit", "Maven", "Data Structures",
    "Algorithms", "OOP", "Linux", "Agile", "Communication", "Problem Solving",
    "MongoDB", "Redis", "Kubernetes", "CI/CD", "Statistics", "Data Analysis"
]
STATUSES = {"Saved", "Applied", "Assessment", "Interview", "Offer", "Rejected", "Withdrawn"}


def json_list(value):
    import json
    try:
        return json.loads(value or "[]")
    except (ValueError, TypeError):
        return []


def analysis_to_dict(item):
    return {
        "id": item.id,
        "createdAt": item.created_at.isoformat() if item.created_at else None,
        "fileName": item.file_name,
        "jobDescription": item.job_description,
        "matchScore": item.match_score,
        "resumeSkills": json_list(item.resume_skills_json),
        "jobSkills": json_list(item.job_skills_json),
        "matchingSkills": json_list(item.matching_skills_json),
        "missingSkills": json_list(item.missing_skills_json),
        "recommendations": json_list(item.recommendations_json),
        "extractedCharacterCount": item.extracted_character_count,
        "detectedJobSkillCount": item.detected_job_skill_count,
        "analysisMethod": item.analysis_method,
    }


def application_to_dict(item):
    return {
        "id": item.id,
        "jobTitle": item.job_title,
        "company": item.company,
        "jobUrl": item.job_url,
        "status": item.status,
        "appliedDate": item.applied_date.isoformat() if item.applied_date else None,
        "notes": item.notes,
        "createdAt": item.created_at.isoformat() if item.created_at else None,
    }


def extract_resume_text(file_storage):
    filename = secure_filename(file_storage.filename or "")
    suffix = Path(filename).suffix.lower()
    if suffix not in {".pdf", ".docx"}:
        raise ValueError("Please upload a PDF or DOCX resume.")
    file_storage.stream.seek(0)
    if suffix == ".pdf":
        if PdfReader is None:
            raise RuntimeError("PDF support is not installed. Run: pip install -r requirements.txt")
        try:
            reader = PdfReader(file_storage.stream)
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as exc:
            raise ValueError("Could not read this PDF. Please try another PDF file.") from exc
    else:
        if Document is None:
            raise RuntimeError("DOCX support is not installed. Run: pip install -r requirements.txt")
        try:
            doc = Document(file_storage.stream)
            text = "\n".join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                for row in table.rows:
                    text += "\n" + " ".join(cell.text for cell in row.cells)
        except Exception as exc:
            raise ValueError("Could not read this DOCX file. Please check that it is valid.") from exc
    if not text.strip():
        raise ValueError("The resume contains no readable text. Scanned image-only PDFs may need OCR.")
    return filename or "Resume", text


def find_skills(text):
    found = []
    lower = text.lower()
    for skill in SKILL_LIBRARY:
        pattern = r"(?<![a-z0-9])" + re.escape(skill.lower()) + r"(?![a-z0-9])"
        if re.search(pattern, lower):
            found.append(skill)
    return found


def build_recommendations(job_skills, matching, missing, resume_skills):
    result = []
    if not job_skills:
        result.append("No supported technical skills were detected in the job description. Add clear skill requirements for a more useful comparison.")
    elif not missing:
        result.append("Your resume mentions all detected skills from this job description. Verify that your experience demonstrates each skill.")
    else:
        result.append("Review these job requirements and build practical experience with the missing skills: " + ", ".join(missing) + ".")
    if matching:
        result.append("Highlight your experience with " + ", ".join(matching) + " in your skills and project sections.")
    else:
        result.append("Add relevant projects or work experience where they truthfully demonstrate the role's requirements.")
    if not resume_skills:
        result.append("No supported skills were detected in your resume text. Check the extracted text and use clear skill names.")
    result.append("Tailor project descriptions to the role and include measurable results only when you can substantiate them.")
    return result


def analyze_text(resume_text, job_description):
    if not resume_text.strip():
        raise ValueError("The resume contains no readable text.")
    if not job_description.strip():
        raise ValueError("Please provide a job description.")
    resume_skills = find_skills(resume_text)
    job_skills = find_skills(job_description)
    matching = [skill for skill in job_skills if skill in resume_skills]
    missing = [skill for skill in job_skills if skill not in resume_skills]
    score = round(len(matching) * 100 / len(job_skills)) if job_skills else 0
    return {
        "matchScore": score,
        "resumeSkills": resume_skills,
        "jobSkills": job_skills,
        "matchingSkills": matching,
        "missingSkills": missing,
        "recommendations": build_recommendations(job_skills, matching, missing, resume_skills),
        "extractedCharacterCount": len(resume_text),
        "detectedJobSkillCount": len(job_skills),
        "analysisMethod": "Rule-based skill comparison (no paid API)",
    }


def skill_from_payload(payload):
    return [str(x).strip() for x in (payload or []) if str(x).strip()]


@app.get("/api/health")
def health():
    return jsonify({"status": "UP", "message": "Backend connected", "language": "Python"})


@app.get("/api/ai/status")
def ai_status():
    return jsonify({"configured": False, "provider": "local-rule-based", "message": "No paid AI API is required."})


@app.get("/api/analyses")
def get_analyses():
    with SessionLocal() as db:
        rows = db.scalars(select(SavedAnalysis).order_by(desc(SavedAnalysis.created_at))).all()
        return jsonify([analysis_to_dict(row) for row in rows])


@app.post("/api/analyze")
def analyze_resume():
    if "resume" not in request.files:
        return jsonify({"error": "Please upload your resume first."}), 400
    job_description = (request.form.get("jobDescription") or "").strip()
    try:
        filename, resume_text = extract_resume_text(request.files["resume"])
        result = analyze_text(resume_text, job_description)
        import json
        with SessionLocal() as db:
            row = SavedAnalysis(
                file_name=filename,
                job_description=job_description,
                match_score=result["matchScore"],
                resume_skills_json=json.dumps(result["resumeSkills"]),
                job_skills_json=json.dumps(result["jobSkills"]),
                matching_skills_json=json.dumps(result["matchingSkills"]),
                missing_skills_json=json.dumps(result["missingSkills"]),
                recommendations_json=json.dumps(result["recommendations"]),
                extracted_character_count=result["extractedCharacterCount"],
                detected_job_skill_count=result["detectedJobSkillCount"],
                analysis_method=result["analysisMethod"],
            )
            db.add(row)
            db.commit()
            db.refresh(row)
            return jsonify(analysis_to_dict(row))
    except (ValueError, RuntimeError) as exc:
        return jsonify({"error": str(exc)}), 400
    except Exception:
        app.logger.exception("Resume analysis failed")
        return jsonify({"error": "Analysis failed. Check backend logs and database configuration."}), 500


@app.post("/api/ai/recommendations")
def recommendations():
    payload = request.get_json(silent=True) or {}
    job_skills = skill_from_payload(payload.get("jobSkills"))
    resume_skills = skill_from_payload(payload.get("resumeSkills"))
    matching = skill_from_payload(payload.get("matchingSkills"))
    missing = skill_from_payload(payload.get("missingSkills"))
    result = build_recommendations(job_skills, matching, missing, resume_skills)
    if payload.get("matchScore") is not None:
        result.insert(0, f"Current rule-based match score: {int(payload.get('matchScore') or 0)}%. This is a skill-keyword estimate, not a hiring prediction.")
    result.extend([
        "Prioritize one or two missing skills that appear most often in the role description.",
        "Prepare a truthful project example for each important matching skill, including your own contribution and how you tested it.",
        "Do not add a skill to your resume until you can explain and demonstrate it."
    ])
    return jsonify({"recommendations": "\n".join(f"• {line}" for line in result)})


def make_question(category, question, suggested):
    return {"category": category, "question": question, "suggestedAnswer": suggested}


@app.post("/api/interview/questions")
def interview_questions():
    payload = request.get_json(silent=True) or {}
    job_description = (payload.get("jobDescription") or "").strip()
    if not job_description:
        return jsonify({"error": "Enter a job description to generate interview questions."}), 400
    skills = skill_from_payload(payload.get("skills"))
    for skill in find_skills(job_description):
        if skill not in skills:
            skills.append(skill)
    questions = [
        make_question("General", "Tell me about yourself and why you are interested in this role.",
                      "Give a 60–90 second summary: your education or experience, two relevant skills, one project with your contribution, and why those strengths fit this role."),
        make_question("Projects", "Walk me through a project relevant to this position.",
                      "Explain the problem, your role, technology choices, one challenge, how you tested it, and the result. Be specific about what you personally implemented."),
        make_question("Problem solving", "Describe a technical problem you faced and how you solved it.",
                      "Use the STAR structure: situation, task, action, result. Explain how you investigated the issue and what you learned."),
    ]
    for skill in skills[:8]:
        low = skill.lower()
        if "sql" in low or "mysql" in low or "database" in low:
            q = f"How would you design and query a relational database using {skill}?"
            hint = "Discuss tables, primary and foreign keys, relationships, constraints, indexes, joins, and how you would validate query results."
        elif "react" in low or "javascript" in low:
            q = f"How would you build and debug an interactive UI using {skill}?"
            hint = "Describe components, state and props, event handling, loading/error states, and how you would inspect browser console/network errors."
        elif any(x in low for x in ["python", "java", "flask", "spring", "fastapi", "oop"]):
            q = f"How have you used {skill} to structure or implement an application?"
            hint = "Explain the relevant concepts, the design you chose, one trade-off, and a concrete example from your own code."
        elif "machine learning" in low or "deep learning" in low:
            q = "How would you evaluate a machine-learning model for a practical problem?"
            hint = "Explain train/validation/test splits, an appropriate metric, class imbalance, overfitting checks, and limitations."
        else:
            q = f"Describe a practical task where you would use {skill} and how you would approach it."
            hint = "Define the problem, explain the core concept, outline steps, mention testing, and give an example you can honestly support."
        questions.append(make_question("Technical · " + skill, q, hint))
    questions.extend([
        make_question("Behavioural", "What would you do if you did not know how to solve a task assigned to you?",
                      "Clarify the expected outcome, break the task into smaller parts, research trusted documentation, test a small solution, and ask a focused question when blocked."),
        make_question("Role fit", "What would you like to learn or achieve in your first 90 days?",
                      "Learn the codebase and team workflow, deliver a small task, seek feedback, and gradually own a measurable feature.")
    ])
    return jsonify(questions)


@app.post("/api/ai/interview-feedback")
def interview_feedback():
    payload = request.get_json(silent=True) or {}
    answer = (payload.get("answer") or "").strip()
    question = (payload.get("question") or "").strip()
    if not answer:
        return jsonify({"error": "Write your answer before requesting feedback."}), 400

    words = re.findall(r"\b[\w'-]+\b", answer)
    lower = answer.lower()
    strengths = []
    improvements = []

    # Transparent rule-based rubric: five criteria, 20 points each.
    detail_score = 20 if len(words) >= 40 else (10 if len(words) >= 25 else 0)
    action_terms = ["i built", "i developed", "i implemented", "i tested", "my role", "i designed", "i created", "i fixed", "i led", "i analyzed", "i analysed"]
    action_score = 20 if any(term in lower for term in action_terms) else 0
    example_terms = ["example", "for instance", "project", "when i", "during my", "in my"]
    example_score = 20 if any(term in lower for term in example_terms) else 0
    reasoning_terms = ["because", "therefore", "so that", "which helped", "in order to", "the reason", "approach"]
    reasoning_score = 20 if any(term in lower for term in reasoning_terms) else 0
    result_terms = ["result", "improved", "increased", "reduced", "learned", "learnt", "completed", "delivered", "successfully", "feedback"]
    result_score = 20 if any(term in lower for term in result_terms) else 0

    criteria = [
        {"name": "Answer detail", "score": detail_score, "maxScore": 20},
        {"name": "Personal contribution", "score": action_score, "maxScore": 20},
        {"name": "Specific example", "score": example_score, "maxScore": 20},
        {"name": "Reasoning", "score": reasoning_score, "maxScore": 20},
        {"name": "Result or learning", "score": result_score, "maxScore": 20},
    ]
    score = sum(item["score"] for item in criteria)

    if detail_score == 20:
        strengths.append("Your answer has enough detail to evaluate.")
    elif detail_score == 10:
        improvements.append("Add a little more detail; aim for roughly 40–100 words for most interview answers.")
    else:
        improvements.append("Expand your answer; aim for roughly 40–100 words for most interview answers.")
    if action_score:
        strengths.append("You describe your own actions rather than only general team activity.")
    else:
        improvements.append("Explain what you personally did using specific actions such as designed, implemented, debugged, or tested.")
    if example_score:
        strengths.append("You connect your answer to a project or real situation.")
    else:
        improvements.append("Include one concrete project or real situation as evidence.")
    if reasoning_score:
        strengths.append("You explain some of your reasoning or approach.")
    else:
        improvements.append("Explain why you chose your approach, using a phrase such as ‘because’ or ‘so that’.")
    if result_score:
        strengths.append("You mention an outcome, result, or lesson learned.")
    else:
        improvements.append("Finish with the outcome, what improved, or what you learned.")
    if len(words) > 150:
        improvements.append("Try shortening the answer and leading with the main point.")

    feedback = "Practice feedback (rule-based; not an AI hiring assessment)\n\nStrengths:\n"
    feedback += "\n".join("• " + item for item in strengths) or "• Your answer is ready for more specific detail."
    if improvements:
        feedback += "\n\nTo improve:\n" + "\n".join("• " + item for item in improvements)
    feedback += "\n\nSuggested structure: direct answer → specific example → your actions → result or learning."
    return jsonify({"score": score, "maxScore": 100, "criteria": criteria, "feedback": feedback, "scoringNote": "This is a simple keyword-based practice rubric, not an assessment of hiring potential."})

@app.get("/api/applications")
def list_applications():
    with SessionLocal() as db:
        rows = db.scalars(select(JobApplication).order_by(desc(JobApplication.created_at))).all()
        return jsonify([application_to_dict(row) for row in rows])


@app.post("/api/applications")
def create_application():
    payload = request.get_json(silent=True) or {}
    title = (payload.get("jobTitle") or "").strip()
    company = (payload.get("company") or "").strip()
    status = (payload.get("status") or "Saved").strip()
    if not title or not company:
        return jsonify({"error": "Job title and company are required."}), 400
    if status not in STATUSES:
        return jsonify({"error": "Invalid application status."}), 400
    try:
        applied_date = date.fromisoformat(payload["appliedDate"]) if payload.get("appliedDate") else None
    except ValueError:
        return jsonify({"error": "Applied date must use YYYY-MM-DD format."}), 400
    with SessionLocal() as db:
        row = JobApplication(
            job_title=title, company=company,
            job_url=(payload.get("jobUrl") or "").strip() or None,
            status=status, applied_date=applied_date,
            notes=(payload.get("notes") or "").strip() or None,
        )
        db.add(row)
        db.commit()
        db.refresh(row)
        return jsonify(application_to_dict(row)), 201


@app.patch("/api/applications/<int:application_id>/status")
def update_application_status(application_id):
    payload = request.get_json(silent=True) or {}
    status = (payload.get("status") or "").strip()
    if status not in STATUSES:
        return jsonify({"error": "Invalid application status."}), 400
    with SessionLocal() as db:
        row = db.get(JobApplication, application_id)
        if not row:
            return jsonify({"error": "Application not found."}), 404
        row.status = status
        db.commit()
        db.refresh(row)
        return jsonify(application_to_dict(row))


@app.delete("/api/applications/<int:application_id>")
def delete_application(application_id):
    with SessionLocal() as db:
        row = db.get(JobApplication, application_id)
        if not row:
            return "", 204
        db.delete(row)
        db.commit()
        return "", 204


@app.errorhandler(413)
def too_large(_error):
    return jsonify({"error": "Your resume must be 10 MB or smaller."}), 413


@app.errorhandler(Exception)
def unhandled_error(error):
    # Preserve useful HTTP errors while avoiding leaking tracebacks to the browser.
    from werkzeug.exceptions import HTTPException
    if isinstance(error, HTTPException):
        return jsonify({"error": error.description}), error.code
    app.logger.exception("Unhandled backend error")
    return jsonify({"error": "Unexpected server error. Check the backend terminal for details."}), 500


if __name__ == "__main__":
    port = int(os.getenv("PORT", "8080"))
    app.run(host="127.0.0.1", port=port, debug=True)
